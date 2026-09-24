from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from bot.audit_log import AuditLog
from bot.classifier import ReceiptClassifier
from bot.controller import ReceiptBotController
from bot.models import Direction, DocumentType, ProposedTransaction, ReceiptExtraction
from bot.workbook_writer import WorkbookWriter


class FakeOCRClient:
    def __init__(self, extraction: ReceiptExtraction) -> None:
        self.extraction = extraction

    def extract(self, file_path: Path, mime_type: str) -> ReceiptExtraction:
        return self.extraction


def create_sample_workbook(path: Path) -> None:
    workbook = Workbook()
    raw = workbook.active
    raw.title = "Transactions Raw"
    raw["B2"] = "#"
    raw["C2"] = "Date"
    raw["D2"] = "Direction"
    raw["E2"] = "Account"
    raw["F2"] = "Amount (ZAR)"
    raw["G2"] = "Notes"
    raw["B3"] = 1
    raw["C3"] = date(2026, 6, 1)
    raw["C3"].number_format = "yyyy-mm-dd"
    raw["D3"] = "Expense"
    raw["E3"] = "Petrol"
    raw["F3"] = 100.0
    raw["G3"] = "Initial petrol"
    for column in range(2, 8):
        raw.cell(4, column)._style = raw.cell(3, column)._style

    trial = workbook.create_sheet("FY26 Trail Balance")
    trial["A1"] = "Expense total"
    trial["B1"] = '=SUMIFS(\'Transactions Raw\'!$F$3:$F$3,\'Transactions Raw\'!$D$3:$D$3,"Expense")'
    workbook.save(path)


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    path = tmp_path / "accounting.xlsx"
    create_sample_workbook(path)
    return path


@pytest.fixture
def workbook_writer(tmp_path: Path, workbook_path: Path) -> WorkbookWriter:
    return WorkbookWriter(
        workbook_path=workbook_path,
        output_dir=tmp_path / "outputs",
        backup_dir=tmp_path / "backups",
        backup_retention_count=20,
        export_filename_prefix="Jess_Lombard_Junior_Dance_Centre_Accounting",
    )


@pytest.fixture
def extraction() -> ReceiptExtraction:
    return ReceiptExtraction(
        document_type=DocumentType.RECEIPT,
        merchant_name="Shell",
        transaction_date=date(2026, 6, 10),
        total_amount_zar=Decimal("642.30"),
        currency="ZAR",
        line_items_summary="Fuel",
        direction=Direction.EXPENSE,
        account=None,
        notes="Shell - petrol",
        confidence=0.91,
        needs_user_review=False,
        reasoning_summary="Fuel slip",
    )


@pytest.fixture
def controller(tmp_path: Path, workbook_writer: WorkbookWriter, extraction: ReceiptExtraction) -> ReceiptBotController:
    workbook_writer.output_dir.mkdir(parents=True, exist_ok=True)
    workbook_writer.backup_dir.mkdir(parents=True, exist_ok=True)
    audit_log = AuditLog(tmp_path / "audit.sqlite3")
    return ReceiptBotController(
        ocr_client=FakeOCRClient(extraction),
        classifier=ReceiptClassifier(),
        workbook_writer=workbook_writer,
        audit_log=audit_log,
        authorized_user_ids={12345, 67890},
        default_dry_run=False,
        default_batch_mode=False,
        rate_limit_window_seconds=60,
        rate_limit_max_uploads=8,
    )


def load_raw_sheet(path: Path):
    return load_workbook(path)["Transactions Raw"]


def load_trial_sheet(path: Path):
    return load_workbook(path)["FY26 Trail Balance"]
