from __future__ import annotations

from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from bot.models import Direction, DocumentType, ProposedTransaction


def make_transaction() -> ProposedTransaction:
    return ProposedTransaction(
        transaction_date=date(2026, 6, 10),
        direction=Direction.EXPENSE,
        account="Petrol",
        amount=Decimal("642.30"),
        notes="Shell - petrol",
        confidence=0.91,
        needs_user_review=False,
        reasoning_summary="Fuel slip",
        merchant_name="Shell",
        document_type=DocumentType.RECEIPT,
    )


def test_append_writes_to_correct_row(workbook_writer) -> None:
    result = workbook_writer.append_transaction(make_transaction())
    raw = load_workbook(workbook_writer.workbook_path)["Transactions Raw"]
    assert result.workbook_row_number == 4
    assert raw["C4"].value.date().isoformat() == "2026-06-10"
    assert raw["E4"].value == "Petrol"
    assert raw["F4"].value == 642.3


def test_append_preserves_headers(workbook_writer) -> None:
    workbook_writer.append_transaction(make_transaction())
    raw = load_workbook(workbook_writer.workbook_path)["Transactions Raw"]
    assert raw["B2"].value == "#"
    assert raw["G2"].value == "Notes"


def test_trial_balance_formula_expands_to_new_row(workbook_writer) -> None:
    workbook_writer.append_transaction(make_transaction())
    trial = load_workbook(workbook_writer.workbook_path, data_only=False)["FY26 Trail Balance"]
    assert "$F$3:$F$4" in trial["B1"].value
    assert "$D$3:$D$4" in trial["B1"].value


def test_trial_balance_formula_catches_up_when_existing_range_is_stale(workbook_path, workbook_writer) -> None:
    workbook = load_workbook(workbook_path)
    raw = workbook["Transactions Raw"]
    raw["B4"] = 2
    raw["C4"] = date(2026, 6, 2)
    raw["C4"].number_format = "yyyy-mm-dd"
    raw["D4"] = "Expense"
    raw["E4"] = "Petrol"
    raw["F4"] = 120.0
    raw["G4"] = "Second petrol"
    raw["B5"] = 3
    raw["C5"] = date(2026, 6, 3)
    raw["C5"].number_format = "yyyy-mm-dd"
    raw["D5"] = "Expense"
    raw["E5"] = "Petrol"
    raw["F5"] = 130.0
    raw["G5"] = "Third petrol"
    trial = workbook["FY26 Trail Balance"]
    trial["B1"] = '=SUMIFS(\'Transactions Raw\'!$F$3:$F$3,\'Transactions Raw\'!$D$3:$D$3,"Expense")'
    workbook.save(workbook_path)

    workbook_writer.append_transaction(make_transaction())

    refreshed_trial = load_workbook(workbook_writer.workbook_path, data_only=False)["FY26 Trail Balance"]
    assert "$F$3:$F$6" in refreshed_trial["B1"].value
    assert "$D$3:$D$6" in refreshed_trial["B1"].value


def test_backup_created_before_write(workbook_writer) -> None:
    workbook_writer.append_transaction(make_transaction())
    backups = list(workbook_writer.backup_dir.glob("BACKUP_*.xlsx"))
    assert backups


def test_preview_does_not_overwrite_production(workbook_writer) -> None:
    workbook_writer.create_preview(make_transaction())
    raw = load_workbook(workbook_writer.workbook_path)["Transactions Raw"]
    assert raw["C4"].value is None


def test_undo_clears_last_written_row(workbook_writer) -> None:
    result = workbook_writer.append_transaction(make_transaction())
    workbook_writer.undo_last_transaction(
        workbook_row_number=result.workbook_row_number,
        expected_values={
            "Date": "2026-06-10",
            "Direction": "Expense",
            "Account": "Petrol",
            "Amount": "642.30",
            "Notes": "Shell - petrol",
        },
    )
    raw = load_workbook(workbook_writer.workbook_path)["Transactions Raw"]
    assert raw["C4"].value is None
