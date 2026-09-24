from __future__ import annotations

import re
import shutil
from copy import copy
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import ProposedTransaction, WriteResult

RAW_SHEET = "Transactions Raw"
TRIAL_BALANCE_SHEET = "FY26 Trail Balance"
HEADER_MAP = {"B2": "#", "C2": "Date", "D2": "Direction", "E2": "Account", "F2": "Amount (ZAR)", "G2": "Notes"}
FORMULA_RANGE_PATTERN = re.compile(
    r"(?P<prefix>'?Transactions Raw'?!\$?[A-Z]{1,3}\$?)(?P<start>\d+)(?P<middle>:\$?[A-Z]{1,3}\$?)(?P<end>\d+)"
)


class WorkbookWriter:
    def __init__(self, *, workbook_path: Path, output_dir: Path, backup_dir: Path, backup_retention_count: int, export_filename_prefix: str) -> None:
        self.workbook_path = workbook_path
        self.output_dir = output_dir
        self.backup_dir = backup_dir
        self.backup_retention_count = backup_retention_count
        self.export_filename_prefix = export_filename_prefix
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.backup_dir.mkdir(parents=True, exist_ok=True)

    def create_backup(self) -> Path:
        timestamp = datetime.now(UTC).strftime("%Y-%m-%d_%H%M%S")
        backup_path = self.backup_dir / f"BACKUP_{self.export_filename_prefix}_{timestamp}.xlsx"
        shutil.copy2(self.workbook_path, backup_path)
        self._prune_old_backups()
        return backup_path

    def append_transaction(self, transaction: ProposedTransaction) -> WriteResult:
        backup_path = self.create_backup()
        workbook = load_workbook(self.workbook_path)
        row_number = self._append_to_workbook(workbook, transaction)
        workbook.save(self.workbook_path)
        self.validate_workbook(self.workbook_path, expected_transaction=transaction, expected_row=row_number)
        return WriteResult(workbook_path=self.workbook_path, backup_path=backup_path, output_path=None, workbook_row_number=row_number)

    def create_preview(self, transaction: ProposedTransaction) -> WriteResult:
        timestamp = datetime.now(UTC).strftime("%Y-%m-%d_%H%M")
        preview_path = self.output_dir / f"PREVIEW_{self.export_filename_prefix}_{timestamp}.xlsx"
        shutil.copy2(self.workbook_path, preview_path)
        workbook = load_workbook(preview_path)
        row_number = self._append_to_workbook(workbook, transaction)
        workbook.save(preview_path)
        self.validate_workbook(preview_path, expected_transaction=transaction, expected_row=row_number)
        return WriteResult(
            workbook_path=preview_path,
            backup_path=preview_path,
            output_path=preview_path,
            workbook_row_number=row_number,
            preview=True,
        )

    def export_workbook(self, *, source_path: Path | None = None) -> Path:
        source = source_path or self.workbook_path
        timestamp = datetime.now(UTC).strftime("%Y-%m-%d_%H%M")
        export_path = self.output_dir / f"{self.export_filename_prefix}_updated_{timestamp}.xlsx"
        shutil.copy2(source, export_path)
        self.validate_workbook(export_path)
        return export_path

    def undo_last_transaction(self, *, workbook_row_number: int, expected_values: dict[str, str]) -> Path:
        backup_path = self.create_backup()
        workbook = load_workbook(self.workbook_path)
        sheet = self._get_raw_sheet(workbook)
        last_row = self._find_last_transaction_row(sheet)
        if workbook_row_number != last_row:
            raise ValueError("Undo is only safe for the most recent transaction row")
        if not self._row_matches(sheet, workbook_row_number, expected_values):
            raise ValueError("Workbook row no longer matches the audit log; refusing automatic undo")
        for column in range(2, 8):
            sheet.cell(workbook_row_number, column).value = None
        workbook.save(self.workbook_path)
        return backup_path

    def validate_workbook(self, workbook_path: Path, *, expected_transaction: ProposedTransaction | None = None, expected_row: int | None = None) -> None:
        workbook = load_workbook(workbook_path)
        raw_sheet = self._get_raw_sheet(workbook)
        if TRIAL_BALANCE_SHEET not in workbook.sheetnames:
            raise ValueError("FY26 Trail Balance sheet missing")
        for cell, expected in HEADER_MAP.items():
            if raw_sheet[cell].value != expected:
                raise ValueError(f"Header {cell} is missing or changed")
        if expected_transaction and expected_row:
            if raw_sheet.cell(expected_row, 3).value is None:
                raise ValueError("Written transaction row is missing")
            if not isinstance(raw_sheet.cell(expected_row, 6).value, (int, float)):
                raise ValueError("Amount must be stored as numeric")
            if raw_sheet.cell(expected_row, 4).value not in {"Expense", "Income"}:
                raise ValueError("Direction must be Expense or Income")
            if raw_sheet.cell(expected_row, 5).value != expected_transaction.account:
                raise ValueError("Account mismatch after write")
        trial_sheet = workbook[TRIAL_BALANCE_SHEET]
        if not any(cell.data_type == "f" for row in trial_sheet.iter_rows() for cell in row if isinstance(cell, Cell)):
            raise ValueError("Trial Balance formulas are missing")

    def _append_to_workbook(self, workbook: Workbook, transaction: ProposedTransaction) -> int:
        raw_sheet = self._get_raw_sheet(workbook)
        last_row = self._find_last_transaction_row(raw_sheet)
        row_number = max(3, last_row + 1)

        self._copy_row_style(raw_sheet, last_row, row_number)
        raw_sheet.cell(row_number, 2).value = self._next_sequence_number(raw_sheet, last_row)
        raw_sheet.cell(row_number, 3).value = transaction.transaction_date
        raw_sheet.cell(row_number, 3).number_format = "yyyy-mm-dd"
        raw_sheet.cell(row_number, 4).value = str(transaction.direction)
        raw_sheet.cell(row_number, 5).value = transaction.account
        raw_sheet.cell(row_number, 6).value = float(transaction.amount)
        raw_sheet.cell(row_number, 7).value = transaction.notes

        self._expand_formula_ranges(workbook, last_row=last_row, new_row=row_number)
        return row_number

    def _get_raw_sheet(self, workbook: Workbook) -> Worksheet:
        if RAW_SHEET not in workbook.sheetnames:
            raise ValueError("Transactions Raw sheet missing")
        return workbook[RAW_SHEET]

    def _find_last_transaction_row(self, sheet: Worksheet) -> int:
        last_row = 2
        for row in range(3, sheet.max_row + 1):
            values = [sheet.cell(row, col).value for col in range(3, 8)]
            if any(value not in (None, "") for value in values):
                last_row = row
        return last_row

    def _next_sequence_number(self, sheet: Worksheet, last_row: int) -> int:
        for row in range(last_row, 2, -1):
            value = sheet.cell(row, 2).value
            if isinstance(value, int):
                return value + 1
            if isinstance(value, float):
                return int(value) + 1
        return 1

    def _copy_row_style(self, sheet: Worksheet, source_row: int, target_row: int) -> None:
        if source_row < 3 or source_row == target_row:
            return
        for column in range(2, 8):
            source = sheet.cell(source_row, column)
            target = sheet.cell(target_row, column)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.font:
                target.font = copy(source.font)
            if source.fill:
                target.fill = copy(source.fill)
            if source.border:
                target.border = copy(source.border)
            if source.alignment:
                target.alignment = copy(source.alignment)
            if source.protection:
                target.protection = copy(source.protection)

    def _expand_formula_ranges(self, workbook: Workbook, *, last_row: int, new_row: int) -> None:
        if new_row <= last_row:
            return
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type != "f" or not isinstance(cell.value, str):
                        continue
                    cell.value = FORMULA_RANGE_PATTERN.sub(
                        lambda match: self._replace_formula_range(match, last_row=last_row, new_row=new_row),
                        cell.value,
                    )

    def _replace_formula_range(self, match: re.Match[str], *, last_row: int, new_row: int) -> str:
        start = int(match.group("start"))
        end = int(match.group("end"))
        # Some live workbooks already have formulas lagging behind the actual
        # last populated transaction row. Any fixed range ending before the new
        # append row must be extended so Trial Balance keeps including new data.
        if start > end or end >= new_row:
            return match.group(0)
        return f"{match.group('prefix')}{start}{match.group('middle')}{new_row}"

    def _row_matches(self, sheet: Worksheet, row_number: int, expected_values: dict[str, str]) -> bool:
        date_value = sheet.cell(row_number, 3).value
        if isinstance(date_value, datetime):
            actual_date = date_value.date().isoformat()
        elif isinstance(date_value, date):
            actual_date = date_value.isoformat()
        else:
            actual_date = str(date_value)
        actual = {
            "Date": actual_date,
            "Direction": str(sheet.cell(row_number, 4).value),
            "Account": str(sheet.cell(row_number, 5).value),
            "Amount": f"{float(sheet.cell(row_number, 6).value):.2f}" if sheet.cell(row_number, 6).value is not None else "",
            "Notes": str(sheet.cell(row_number, 7).value),
        }
        return actual == expected_values

    def _prune_old_backups(self) -> None:
        backups = sorted(self.backup_dir.glob(f"BACKUP_{self.export_filename_prefix}_*.xlsx"))
        while len(backups) > self.backup_retention_count:
            oldest = backups.pop(0)
            oldest.unlink(missing_ok=True)
